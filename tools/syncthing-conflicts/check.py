#!/usr/bin/env python3
"""
Разбор конфликтов Syncthing в базе Obsidian. Запуск — systemd-таймер раз в 10 минут.

1. `*.sync-conflict-*`:
   - копия совпадает с оригиналом или это служебный мусор — удаляется;
   - оригинала нет — копия встаёт на его место;
   - содержимое разное — копия уезжает в архив, в Telegram отчёт с разницей.
2. `<имя> (сервер).<ext>` от office_guard: файл закрыли без правок —
   копия встаёт на место, старый оригинал в архив; правили — просим свести.
3. Файлы-владельцы `~$…`: старше суток — предупреждение, старше трёх дней —
   в архив (office_guard их уже не учитывает).
4. Архив старше 60 дней чистится.

Архив: /root/Obsidian/.stversions/_conflicts/ (в игноре Syncthing, живёт только
на сервере). Отчёт — только если что-то произошло. --dry-run ничего не меняет
и печатает отчёт в stdout.
"""

import difflib
import filecmp
import hashlib
import json
import os
import re
import shutil
import subprocess
import sys
import time
import xml.etree.ElementTree as ET
from pathlib import Path

sys.path.insert(0, "/usr/local/lib/python3.12/dist-packages")
from office_guard import SIDE_SUFFIX, STALE_LOCK_S, find_lock  # noqa: E402

VAULT = Path(os.environ.get("OBSIDIAN_VAULT", "/root/Obsidian"))  # переопределение — для тестов
ARCHIVE = VAULT / ".stversions" / "_conflicts"
STATE = Path(os.environ.get("CONFLICTS_STATE", "/var/lib/syncthing-conflicts/state.json"))
ST_CONFIG = Path("/root/.local/state/syncthing/config.xml")
NOTIFY = os.environ.get("CONFLICTS_NOTIFY", "/opt/claude-telegram/notify.sh")
# Куда слать отчёт. Задаётся в unit-файле (drop-in), без него отчёт только печатается в журнал
CHAT_ID = os.environ.get("CONFLICTS_CHAT_ID", "")
THREAD_REPORTS = os.environ.get("CONFLICTS_THREAD", "")  # топик «Отчёты»

SETTLE_S = 120              # свежий файл может ещё докачиваться
WARN_LOCK_S = 24 * 3600
ARCHIVE_TTL_S = 60 * 24 * 3600
TG_LIMIT = 3800
CONFLICT_RE = re.compile(r"\.sync-conflict-\d{8}-\d{6}-([A-Z0-9]{7})")
SKIP_DIRS = {".stversions", ".stfolder", ".trash"}
JUNK_PARTS = {".git"}
JUNK_NAMES = re.compile(r"^(~\$|\.~lock\.)|^workspace(-mobile)?\.json$")

DRY = "--dry-run" in sys.argv


def rel(p: Path) -> str:
    return str(p.relative_to(VAULT))


def age(p: Path) -> float:
    return time.time() - p.stat().st_mtime


def walk(pattern_ok):
    for p in VAULT.rglob("*"):
        parts = p.relative_to(VAULT).parts
        if parts and parts[0] in SKIP_DIRS:
            continue
        if p.is_file() and pattern_ok(p):
            yield p


def device_names() -> dict[str, str]:
    try:
        root = ET.parse(ST_CONFIG).getroot()
    except (OSError, ET.ParseError):
        return {}
    return {d.get("id", "")[:7]: d.get("name") for d in root.iter("device") if d.get("name")}


def to_archive(p: Path) -> Path:
    dest = ARCHIVE / p.relative_to(VAULT)
    if dest.exists():
        dest = dest.with_name(f"{dest.stem}.{int(time.time())}{dest.suffix}")
    if not DRY:
        dest.parent.mkdir(parents=True, exist_ok=True)
        shutil.move(p, dest)
    return dest


# ---------- разница для отчёта ----------

def _xlsx_cells(p: Path) -> list[str]:
    from openpyxl import load_workbook
    wb = load_workbook(p, read_only=True)
    out = []
    for ws in wb.worksheets:
        for row in ws.iter_rows():
            for c in row:
                if c.value is not None and hasattr(c, "coordinate"):
                    out.append(f"{ws.title}!{c.coordinate}: {c.value}")
    return out


def _docx_lines(p: Path) -> list[str]:
    import docx
    d = docx.Document(p)
    lines = [x.text for x in d.paragraphs if x.text.strip()]
    lines += [c.text for t in d.tables for r in t.rows for c in r.cells if c.text.strip()]
    return lines


def _text_lines(p: Path) -> list[str]:
    return p.read_text(encoding="utf-8", errors="replace").splitlines()


READERS = {".xlsx": _xlsx_cells, ".xlsm": _xlsx_cells, ".docx": _docx_lines,
           ".md": _text_lines, ".txt": _text_lines, ".csv": _text_lines,
           ".json": _text_lines, ".canvas": _text_lines}


def diff_summary(kept: Path, lost: Path, legend: str, limit=4) -> str:
    reader = READERS.get(kept.suffix.lower())
    if reader is None:
        return f"   размер: осталось {kept.stat().st_size} Б, в архиве {lost.stat().st_size} Б"
    try:
        a, b = reader(kept), reader(lost)
    except Exception as e:  # битый файл не должен ронять весь разбор
        return f"   разницу не прочитать: {type(e).__name__}"
    changes = [l for l in difflib.unified_diff(a, b, lineterm="", n=0)
               if l[:1] in "+-" and not l.startswith(("+++", "---"))]
    if not changes:
        return "   по содержимому одинаковые (отличаются служебные данные)"
    shown = [f"   ({legend})"] + [f"   {l[0]} {l[1:][:140]}" for l in changes[:limit]]
    more = len(changes) - limit
    return "\n".join(shown + ([f"   … ещё {more}"] if more > 0 else []))


# ---------- 1. конфликты ----------

def handle_conflicts(report, devices):
    for c in walk(lambda p: CONFLICT_RE.search(p.name)):
        if age(c) < SETTLE_S:
            continue
        m = CONFLICT_RE.search(c.name)
        orig = c.with_name(c.name[:m.start()] + c.name[m.end():])
        who = devices.get(m.group(1), m.group(1))
        parts = set(c.relative_to(VAULT).parts)
        if parts & JUNK_PARTS or JUNK_NAMES.search(orig.name):
            DRY or c.unlink()
            report["junk"].append(rel(c))
        elif not orig.exists():
            DRY or c.rename(orig)
            report["restored"].append(f"• {rel(orig)}\n   оригинала не было — копия ({who}) встала на его место")
        elif filecmp.cmp(c, orig, shallow=False):
            DRY or c.unlink()
            report["same"].append(rel(orig))
        else:
            summary = diff_summary(orig, c, "− осталось, + в архиве")
            dest = to_archive(c)
            report["archived"].append(
                f"• {rel(orig)}\n   проиграла версия с «{who}», она в архиве:\n"
                f"   {rel(dest)}\n{summary}")


# ---------- 2. отложенные копии office_guard ----------

def handle_side_copies(report, state):
    for side in walk(lambda p: p.stem.endswith(SIDE_SUFFIX)):
        orig = side.with_name(side.stem[: -len(SIDE_SUFFIX)] + side.suffix)
        if age(side) < SETTLE_S or find_lock(orig):
            continue  # файл ещё открыт на ПК — ждём
        if not orig.exists() or orig.stat().st_mtime <= side.stat().st_mtime:
            if orig.exists():
                to_archive(orig)
            DRY or side.rename(orig)
            state["side"].pop(rel(side), None)
            report["promoted"].append(f"• {rel(orig)}\n   файл закрыли — серверная версия встала на место")
            continue
        key = rel(side)
        mark = hashlib.md5(f"{orig.stat().st_mtime}".encode()).hexdigest()
        if state["side"].get(key) == mark:
            continue  # уже сообщали, с тех пор ничего не поменялось
        state["side"][key] = mark
        report["manual"].append(
            f"• {rel(orig)}\n   файл правили в Excel/Word, пока сервер писал копию «{side.name}».\n"
            f"   Свести вручную: нужное перенести в оригинал, копию удалить.\n"
            + diff_summary(orig, side, "− в оригинале, + у сервера"))


# ---------- 3. зависшие блокировки ----------

def handle_stale_locks(report, state):
    seen = set()
    for lock in walk(lambda p: p.name.startswith("~$")):
        key = rel(lock)
        if age(lock) > STALE_LOCK_S:
            to_archive(lock)
            report["junk"].append(key)
            continue
        seen.add(key)
        if age(lock) > WARN_LOCK_S and key not in state["locks"]:
            state["locks"].append(key)
            report["locks"].append(
                f"• {key}\n   висит больше суток: файл открыт или Excel/Word закрылся с ошибкой. "
                f"Пока он есть, сервер пишет в копию «(сервер)»; через 3 дня уберу сам.")
    state["locks"] = [k for k in state["locks"] if k in seen]


# ---------- 4. чистка архива ----------

def prune_archive():
    if DRY or not ARCHIVE.exists():
        return
    for p in ARCHIVE.rglob("*"):
        if p.is_file() and age(p) > ARCHIVE_TTL_S:
            p.unlink()
    for d in sorted((d for d in ARCHIVE.rglob("*") if d.is_dir()), reverse=True):
        if not any(d.iterdir()):
            d.rmdir()


# ---------- отчёт ----------

SECTIONS = [
    ("manual", "✋ Нужно свести вручную"),
    ("archived", "📦 Конфликт: оставлена победившая версия, вторая в архиве"),
    ("restored", "↩️ Восстановлено из копии"),
    ("promoted", "✅ Отложенная запись сервера встала на место"),
    ("locks", "🔒 Зависшие блокировки"),
]


def render(report) -> str:
    blocks = [f"{title}\n" + "\n".join(report[k]) for k, title in SECTIONS if report[k]]
    quiet = len(report["same"]) + len(report["junk"])
    if quiet:
        blocks.append(f"🧹 Удалено копий-дублей и служебного мусора: {quiet}")
    if not blocks:
        return ""
    text = "🔄 Syncthing · Obsidian\n\n" + "\n\n".join(blocks)
    return text if len(text) <= TG_LIMIT else text[:TG_LIMIT] + "\n…(обрезано)"


def load_state():
    try:
        s = json.loads(STATE.read_text())
    except (OSError, ValueError):
        s = {}
    return {"side": s.get("side", {}), "locks": s.get("locks", [])}


def main():
    state = load_state()
    report = {k: [] for k in ("manual", "archived", "restored", "promoted", "locks", "same", "junk")}
    devices = device_names()
    handle_conflicts(report, devices)
    handle_side_copies(report, state)
    handle_stale_locks(report, state)
    prune_archive()
    text = render(report)
    if DRY:
        print(text or "нечего разбирать")
        return
    STATE.parent.mkdir(parents=True, exist_ok=True)
    STATE.write_text(json.dumps(state, ensure_ascii=False, indent=1))
    if text and not CHAT_ID:
        print("CONFLICTS_CHAT_ID не задан — отчёт не отправлен:\n" + text)
    elif text:
        cmd = [NOTIFY, CHAT_ID] + (["--thread", THREAD_REPORTS] if THREAD_REPORTS else []) + [text]
        subprocess.run(cmd, check=True, stdout=subprocess.DEVNULL)


if __name__ == "__main__":
    main()
